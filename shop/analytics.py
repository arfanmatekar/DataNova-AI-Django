"""Safe file parsing and deterministic, dataset-grounded analytics."""

import csv
import io
import json
import math
import statistics
from datetime import date, datetime

from django.core.exceptions import ValidationError


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 100
SUPPORTED_SUFFIXES = {'.csv', '.xlsx', '.json'}


def _normalize_headers(raw_headers):
    if len(raw_headers) > MAX_COLUMNS:
        raise ValidationError(f'Files can contain at most {MAX_COLUMNS} columns.')
    headers = []
    seen = set()
    for index, value in enumerate(raw_headers[:MAX_COLUMNS], start=1):
        name = str(value).strip() if value is not None else ''
        name = name or f'Column {index}'
        candidate = name
        suffix = 2
        while candidate.casefold() in seen:
            candidate = f'{name} ({suffix})'
            suffix += 1
        seen.add(candidate.casefold())
        headers.append(candidate)
    if not headers:
        raise ValidationError('The file must contain a header row and at least one column.')
    return headers


def _cell_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _finish_dataset(headers, rows):
    if len(headers) > MAX_COLUMNS:
        raise ValidationError(f'Files can contain at most {MAX_COLUMNS} columns.')
    if not rows:
        raise ValidationError('The file must contain at least one data row.')
    if len(rows) > MAX_ROWS:
        raise ValidationError(f'Files can contain at most {MAX_ROWS} data rows.')
    return headers, rows


def _parse_csv(content):
    try:
        text = content.decode('utf-8-sig')
        reader = csv.reader(io.StringIO(text, newline=''))
        raw_headers = next(reader, None)
        if raw_headers is None:
            raise ValidationError('The CSV file is empty.')
        headers = _normalize_headers(raw_headers)
        rows = []
        for line_number, values in enumerate(reader, start=2):
            if not values or not any(value.strip() for value in values):
                continue
            if len(values) > len(headers):
                raise ValidationError(
                    f'CSV row {line_number} has more values than the header.'
                )
            rows.append({
                header: values[index].strip() if index < len(values) else ''
                for index, header in enumerate(headers)
            })
            if len(rows) > MAX_ROWS:
                raise ValidationError(f'Files can contain at most {MAX_ROWS} data rows.')
    except UnicodeDecodeError as exc:
        raise ValidationError('CSV files must use UTF-8 text encoding.') from exc
    except csv.Error as exc:
        raise ValidationError(f'The CSV file could not be read: {exc}') from exc
    return _finish_dataset(headers, rows)


def _parse_json(content):
    try:
        records = json.loads(content.decode('utf-8-sig'))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValidationError('The JSON file is not valid UTF-8 JSON.') from exc
    if isinstance(records, dict) and isinstance(records.get('rows'), list):
        records = records['rows']
    if not isinstance(records, list) or not records:
        raise ValidationError('JSON must contain a non-empty array of objects.')
    if not all(isinstance(record, dict) for record in records):
        raise ValidationError('Every item in the JSON array must be an object.')
    original_keys = list(records[0].keys())
    for record in records[1:]:
        for key in record:
            if key not in original_keys:
                original_keys.append(key)
                if len(original_keys) > MAX_COLUMNS:
                    raise ValidationError(f'Files can contain at most {MAX_COLUMNS} columns.')
    headers = _normalize_headers(original_keys)
    rows = []
    for index, record in enumerate(records[:MAX_ROWS + 1], start=1):
        rows.append({
            headers[column_index]: _cell_value(record.get(key))
            for column_index, key in enumerate(original_keys)
        })
    return _finish_dataset(headers, rows)


def _parse_xlsx(content):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError(
            'Excel support requires OpenPyXL. Install the project dependencies '
            'with "python -m pip install -r requirements.txt".'
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValidationError(f'Files can contain at most {MAX_COLUMNS} columns.')
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise ValidationError(f'Files can contain at most {MAX_ROWS} data rows.')
            iterator = sheet.iter_rows(values_only=True)
            raw_headers = next(iterator, None)
            if raw_headers is None:
                raise ValidationError('The Excel worksheet is empty.')
            headers = _normalize_headers(list(raw_headers))
            rows = []
            for line_number, values in enumerate(iterator, start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                if len(values) > len(headers) and any(
                    value is not None and str(value).strip()
                    for value in values[len(headers):]
                ):
                    raise ValidationError(
                        f'Excel row {line_number} has values beyond the header columns.'
                    )
                rows.append({
                    header: _cell_value(values[index]) if index < len(values) else ''
                    for index, header in enumerate(headers)
                })
                if len(rows) > MAX_ROWS:
                    raise ValidationError(
                        f'Files can contain at most {MAX_ROWS} data rows.'
                    )
        finally:
            workbook.close()
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(
            'The Excel workbook could not be read. Upload a valid .xlsx file.'
        ) from exc
    return _finish_dataset(headers, rows)


def parse_upload(upload):
    filename = upload.name.rsplit('/', 1)[-1].rsplit('\\', 1)[-1]
    suffix = '.' + filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValidationError('Upload a .csv, .xlsx, or .json file.')
    if upload.size > MAX_UPLOAD_BYTES:
        raise ValidationError('Files must be 5 MB or smaller.')
    content = upload.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValidationError('Files must be 5 MB or smaller.')
    if suffix == '.csv':
        columns, rows = _parse_csv(content)
    elif suffix == '.xlsx':
        columns, rows = _parse_xlsx(content)
    else:
        columns, rows = _parse_json(content)
    return {'filename': filename[:255], 'columns': columns, 'rows': rows}


def _as_number(value):
    if isinstance(value, bool) or value is None or value == '':
        return None
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def dataset_summary(columns, rows):
    numeric = {}
    for column in columns:
        values = [
            number for row in rows
            if (number := _as_number(row.get(column))) is not None
        ]
        if values:
            numeric[column] = {
                'count': len(values),
                'sum': round(sum(values), 2),
                'average': round(statistics.fmean(values), 2),
                'minimum': round(min(values), 2),
                'maximum': round(max(values), 2),
            }
    return {
        'row_count': len(rows),
        'column_count': len(columns),
        'numeric': numeric,
    }


def answer_question(question, columns, rows):
    if not isinstance(question, str):
        raise ValidationError('Type a question about your dataset first.')
    question = question.strip()
    if not question:
        raise ValidationError('Type a question about your dataset first.')
    if len(question) > 500:
        raise ValidationError('Questions must be 500 characters or fewer.')
    if not rows:
        raise ValidationError('Upload a dataset before asking a question.')

    summary = dataset_summary(columns, rows)
    numeric = summary['numeric']
    lowered = question.casefold()
    if not numeric:
        return (
            f'Your dataset contains {len(rows):,} rows and {len(columns)} columns. '
            'There are no numeric columns to calculate totals, trends, or anomalies.'
        )

    if any(word in lowered for word in ('anomal', 'outlier', 'unusual')):
        preferred = [
            name for name in numeric
            if name.casefold() in lowered
            or any(term in name.casefold() for term in ('sale', 'revenue', 'price', 'amount'))
        ]
        column = preferred[0] if preferred else next(iter(numeric))
        values = [
            _as_number(row.get(column)) for row in rows
            if _as_number(row.get(column)) is not None
        ]
        median = statistics.median(values)
        deviations = [abs(value - median) for value in values]
        mad = statistics.median(deviations)
        if mad:
            unusual = [value for value in values if abs(value - median) / mad > 3]
        else:
            mean = statistics.fmean(values)
            deviation = statistics.pstdev(values)
            unusual = [
                value for value in values
                if deviation > 0 and abs(value - mean) / deviation > 2.5
            ]
        if unusual:
            examples = ', '.join(f'{value:g}' for value in unusual[:5])
            return (
                f'Found {len(unusual)} potential anomal{"y" if len(unusual) == 1 else "ies"} '
                f'in "{column}" using a robust outlier check. '
                f'Example values: {examples}.'
            )
        return (
            f'No clear outliers stood out in "{column}". '
            f'Its {len(values):,} values range from {min(values):g} to '
            f'{max(values):g}, with an average of {statistics.fmean(values):g}.'
        )

    if any(word in lowered for word in ('trend', 'over time', 'growing', 'growth', 'declining')):
        column = next(
            (name for name in numeric if name.casefold() in lowered),
            next(iter(numeric)),
        )
        values = [
            _as_number(row.get(column)) for row in rows
            if _as_number(row.get(column)) is not None
        ]
        if len(values) < 2:
            return f'There are not enough values in "{column}" to identify a trend.'
        first, last = values[0], values[-1]
        direction = 'upward' if last > first else 'downward' if last < first else 'flat'
        change = (
            f' ({(last - first) / abs(first) * 100:+.1f}%)'
            if first else ''
        )
        return (
            f'The "{column}" series shows an overall {direction} trend across '
            f'{len(values):,} observations, moving from {first:g} to {last:g}{change}.'
        )

    if any(word in lowered for word in ('average', 'mean')):
        matched = next(
            (name for name in numeric if name.casefold() in lowered),
            next(iter(numeric)),
        )
        return (
            f'The average of "{matched}" is '
            f'{summary["numeric"][matched]["average"]:g}, calculated from '
            f'{summary["numeric"][matched]["count"]:,} numeric values.'
        )

    matched = next(
        (name for name in numeric if name.casefold() in lowered),
        next(
            (
                name for name in numeric
                if any(term in name.casefold() for term in ('sale', 'revenue', 'price', 'amount', 'total'))
            ),
            next(iter(numeric)),
        ),
    )
    stats = numeric[matched]
    return (
        f'The total for "{matched}" is {stats["sum"]:g}, '
        f'calculated from {stats["count"]:,} numeric values. '
        f'The average is {stats["average"]:g} and the range is '
        f'{stats["minimum"]:g} to {stats["maximum"]:g}.'
    )
